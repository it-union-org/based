"""
    Desc: transcription, vision, cleanup and formatting pipeline for the notes module
    Creator: Kirosha
"""

from __future__ import annotations

import shutil
import string
import tempfile
import uuid
from contextlib import contextmanager
from pathlib import Path
from typing import Awaitable, Callable

from based.config import config
from based.modules.notes.agents import Summarizer, Transcriber, Vision
from based.modules.notes.schemas import NoteRequest
from based.modules.notes.settings import NotesSettings
from based.modules.notes.utils.__chunker import split_text
from based.utils.__console import log_err, log_step, log_warn, make_progress
from based.utils.__health import HealthState, health_check
from based.utils.__json import write
from based.utils.__llm import ask_text, last_error
from based.utils.__prompts import load_prompt

ProgressCallback = Callable[[str], Awaitable[None]]


class Pipeline:
    def __init__(self, settings: NotesSettings) -> None:
        self.settings = settings
        self.health = HealthState()
        self.transcriber = Transcriber()
        self.vision = Vision()
        self.summarizer = Summarizer()

    @contextmanager
    def __scratch(self, parent: str):
        path = Path(tempfile.gettempdir()) / "based" / parent / str(uuid.uuid4())
        path.mkdir(parents=True, exist_ok=True)
        try:
            yield path
        finally:
            shutil.rmtree(path, ignore_errors=True)

    async def __progress(self, cb: ProgressCallback | None, msg: str) -> None:
        if cb is not None:
            try:
                await cb(msg)
            except Exception:
                pass

    @health_check("create_note")
    async def create_note(
        self,
        request: NoteRequest,
        progress_cb: ProgressCallback | None = None,
    ) -> Path | None:
        if (
            not request.audio_paths
            and not request.image_paths
            and not request.presentation_path
            and not request.raw_text
        ):
            self.health.report_error("nothing to build a note from")
            return log_err("notes: nothing to build a note from")

        out_dir = config.cache_dir / "notes" / str(uuid.uuid4())
        out_dir.mkdir(parents=True, exist_ok=True)

        images = list(request.image_paths)
        transcripts: list[str] = []

        await self.__progress(progress_cb, "Обрабатываю файлы...")

        with self.__scratch("notes") as scratch:
            if request.presentation_path is not None:
                slide_text, slide_images = self.__extract_presentation(
                    request.presentation_path, scratch
                )
                if slide_text:
                    transcripts.append(slide_text)
                images.extend(slide_images)

            if request.audio_paths:
                await self.__progress(progress_cb, "Извлекаю текст из аудио...")
            for audio_path in request.audio_paths:
                log_step(f"notes: transcribing {audio_path.name}")
                text = await self.transcriber.transcribe(audio_path)
                if text is None:
                    self.health.report_error(
                        f"transcription failed for {audio_path.name}: {self.transcriber.last_error}"
                    )
                    continue
                transcripts.append(text)

            if request.raw_text:
                transcripts.insert(0, request.raw_text)

            if images:
                await self.__progress(progress_cb, "Извлекаю текст из фото...")
            image_descriptions = await self.__describe_images(images)

            full = "\n\n".join(transcripts)
            if full:
                await self.__progress(progress_cb, "Генерирую готовый конспект...")
                full = await self.__clean_chunks(full)

            if not full and not image_descriptions:
                self.health.report_error("nothing usable was extracted from the inputs")
                return log_err("notes: nothing usable was extracted from the inputs")

            await self.__progress(progress_cb, "Формирую итоговый документ...")
            return await self.__render_and_save(
                out_dir, request.subject_name, full, "\n\n".join(image_descriptions)
            )

    @health_check("create_note_from_text")
    async def create_note_from_text(
        self,
        text: str,
        subject_name: str,
        progress_cb: ProgressCallback | None = None,
    ) -> Path | None:
        out_dir = config.cache_dir / "notes" / str(uuid.uuid4())
        out_dir.mkdir(parents=True, exist_ok=True)
        await self.__progress(progress_cb, "Генерирую готовый конспект...")
        return await self.__render_and_save(out_dir, subject_name, text, "")

    async def __describe_images(self, images: list[Path]) -> list[str]:
        descriptions: list[str] = []
        if not images:
            return descriptions
        with make_progress() as progress:
            task = progress.add_task("describing images", total=len(images))
            for image_path in images:
                description = await self.vision.describe(image_path)
                if description is None:
                    self.health.report_error(
                        f"vision failed for {image_path.name}: {self.vision.last_error}"
                    )
                else:
                    descriptions.append(description)
                progress.advance(task)
        return descriptions

    async def __clean_chunks(self, text: str) -> str:
        prompt = load_prompt("notes", "clean")
        if prompt is None:
            self.health.report_error("clean prompt not found")
            return text
        chunks = split_text(text, target_size=config.clean_chunk_size)
        log_step(f"cleaning transcript in {len(chunks)} chunk(s)")
        cleaned = []
        with make_progress() as progress:
            task = progress.add_task("cleaning transcript", total=len(chunks))
            for index, chunk in enumerate(chunks):
                result = await ask_text(prompt, chunk)
                if result is None:
                    self.health.report_warning(
                        f"clean failed for chunk {index + 1}/{len(chunks)}: {last_error()}"
                    )
                    cleaned.append(chunk)
                else:
                    cleaned.append(result)
                progress.advance(task)
        return "\n\n".join(cleaned)

    async def __render_and_save(
        self, out_dir: Path, subject_name: str, transcripts: str, image_descriptions: str
    ) -> Path | None:
        user_prompt = load_prompt("notes", "user")
        if user_prompt is None:
            self.health.report_error("user prompt not found")
            return log_err("notes: user prompt not found")

        chunks = (
            split_text(transcripts, target_size=config.format_chunk_size)
            if len(transcripts) > config.format_chunk_size
            else [transcripts]
        )
        if len(chunks) > 1:
            log_step(f"formatting note in {len(chunks)} chunk(s)")

        parts: list[str] = []
        for index, chunk in enumerate(chunks):
            rendered = string.Template(user_prompt).safe_substitute(
                subject_name=subject_name,
                transcripts=chunk,
                image_descriptions=image_descriptions if index == 0 else "",
            )
            content = await self.summarizer.summarize(rendered)
            if content is None:
                cause = self.summarizer.last_error
                self.health.report_error(
                    f"summarization failed for chunk {index + 1}/{len(chunks)}: {cause}"
                )
                return log_err(f"notes: summarization failed: {cause}")
            parts.append(content)

        note_path = out_dir / "note.md"
        note_path.write_text("\n\n".join(parts), encoding="utf-8")
        write(
            out_dir / "stats.json",
            {
                "subject_name": subject_name,
                "transcript_chars": len(transcripts),
                "image_descriptions": len(image_descriptions.split("\n\n")) if image_descriptions else 0,
                "format_chunks": len(chunks),
                "note_path": str(note_path),
            },
        )
        return note_path

    def __extract_presentation(self, path: Path, scratch: Path) -> tuple[str, list[Path]]:
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                return self.__extract_pdf(path, scratch)
            if suffix == ".pptx":
                return self.__extract_pptx(path, scratch)
            log_warn(f"unsupported presentation extension: {suffix}")
        except ImportError as err:
            log_warn(f"missing dependency for {suffix}: {err}")
        except Exception as err:
            self.health.report_warning(f"presentation {path.name} failed: {err.__class__.__name__}: {err}")
            log_warn(f"failed to extract presentation {path.name}: {err.__class__.__name__}: {err}")
        return "", []

    def __extract_pdf(self, path: Path, scratch: Path) -> tuple[str, list[Path]]:
        import pdfplumber
        import pypdfium2 as pdfium

        text_parts: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

        text = "\n".join(text_parts)
        if len(text) >= config.min_text_chars_per_page * page_count:
            return text, []

        images: list[Path] = []
        pdf_doc = pdfium.PdfDocument(str(path))
        for index, page in enumerate(pdf_doc):
            bitmap = page.render(scale=config.presentation_render_scale)
            image_path = scratch / f"page_{index}.png"
            bitmap.to_pil().save(image_path)
            images.append(image_path)
        return text, images

    def __extract_pptx(self, path: Path, scratch: Path) -> tuple[str, list[Path]]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(str(path))
        text_parts: list[str] = []
        images: list[Path] = []
        for slide_index, slide in enumerate(presentation.slides):
            for shape_index, shape in enumerate(slide.shapes):
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs)
                        if text:
                            text_parts.append(text)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image = shape.image
                    image_path = scratch / f"slide_{slide_index}_{shape_index}.{image.ext}"
                    image_path.write_bytes(image.blob)
                    images.append(image_path)
        text = "\n".join(text_parts)
        slide_count = max(1, len(presentation.slides))
        if len(text) >= config.min_text_chars_per_page * slide_count:
            return text, []
        return text, images
